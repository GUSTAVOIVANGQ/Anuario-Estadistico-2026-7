"""Barra común del anuario, tomada de la diapositiva 7 de la plantilla."""

from __future__ import annotations

from copy import deepcopy

from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement


HEADER_DESIGN_VERSION = 1
HEADER_NAMES = (
    "ANUARIO_BARRA_FONDO", "ANUARIO_BARRA_QR_FONDO", "ANUARIO_BARRA_QR",
    "ANUARIO_BARRA_CRT", "ANUARIO_BARRA_TITULO", "ANUARIO_BARRA_CONSULTA",
    "ANUARIO_BARRA_URL",
)
CONSULTATION = "Para consultar información con mayor desagregación, visita"
BIT_LABEL = "BIT CRT · bit.crt.gob.mx/BitWebApp"
BIT_URL = "https://bit.crt.gob.mx/BitWebApp/"


def _text(shape) -> str:
    # Shape.text crea un txBody vacío en figuras sin texto. Leer XML conserva
    # intactos los objetos de la referencia y los fondos al inspeccionarlos.
    paragraphs = shape._element.findall(".//" + qn("a:p"))
    return "\n".join(
        "".join(node.text or "" for node in paragraph.iter(qn("a:t")))
        for paragraph in paragraphs
    )


def _is_header(shape) -> bool:
    # Excluye fondos de página y contenido que sólo comienza cerca del borde.
    return shape.top < 521700 and shape.top + shape.height <= 550000


def _title(slide, number: int) -> str:
    for shape in slide.shapes:
        text = _text(shape).strip()
        if (
            _is_header(shape) and text and "BIT" not in text
            and "consultar" not in text and text != "CRT"
        ):
            return text
    if number == 2:
        return "ÍNDICE"
    raise ValueError(f"Falta el título de sección en la diapositiva {number}.")


def _set_title(element, title: str, number: int) -> None:
    transform = element.find(".//" + qn("a:xfrm"))
    transform.find(qn("a:ext")).set("cx", "5050000")
    body = element.find(qn("p:txBody"))
    paragraph = body.find(qn("a:p"))
    size = 1200 if len(title) > 45 else 1500
    lines = [title]
    if 117 <= number <= 120:
        size = 1200
        lines = [
            "ANEXO I. RESUMEN DE INDICADORES DE",
            "TELECOMUNICACIONES Y RADIODIFUSIÓN",
        ]
    for old in list(body.findall(qn("a:p"))):
        body.remove(old)
    for text in lines:
        copied = deepcopy(paragraph)
        copied.find(".//" + qn("a:t")).text = text
        for properties in copied.iter():
            if properties.tag in (qn("a:rPr"), qn("a:defRPr")):
                properties.set("sz", str(size))
        if len(lines) > 1:
            spacing = OxmlElement("a:lnSpc")
            points = OxmlElement("a:spcPts")
            points.set("val", "1300")
            spacing.append(points)
            copied.find(qn("a:pPr")).insert(0, spacing)
        body.append(copied)
    if len(lines) > 1:
        transform.find(qn("a:off")).set("y", "45720")
        transform.find(qn("a:ext")).set("cy", "365760")


def apply_uniform_headers(presentation) -> dict[str, object]:
    """Aplica la referencia después de construir todas las páginas y figuras.

    La portada, la contraportada y la propia referencia se conservan. El enlace
    se asigna al cuadro completo: un vínculo en el run adopta el color del tema
    de cada sección en PowerPoint, aunque su texto especifique blanco.
    """
    if len(presentation.slides) != 131:
        raise ValueError("La barra común requiere la plantilla de 131 diapositivas.")
    reference = presentation.slides[6]
    originals = [shape for shape in reference.shapes if _is_header(shape)]
    if len(originals) != len(HEADER_NAMES):
        raise ValueError("La diapositiva 7 debe contener los siete objetos de la barra aprobada.")
    if [_text(s) for s in originals[-2:]] != [CONSULTATION, BIT_LABEL]:
        raise ValueError("La diapositiva 7 no contiene la barra BIT CRT aprobada.")
    reference_elements = [deepcopy(shape._element) for shape in originals]
    titles = {n: _title(presentation.slides[n - 1], n) for n in range(2, 131) if n != 7}

    excluded_layouts = {
        presentation.slides[n - 1].slide_layout.part for n in (1, 131)
    }
    layouts = {slide.slide_layout.part: slide.slide_layout for slide in presentation.slides}
    for part, layout in layouts.items():
        if part in excluded_layouts:
            continue
        for shape in list(layout.shapes):
            if _is_header(shape):
                layout.shapes._spTree.remove(shape._element)

    for number in range(2, 131):
        if number == 7:
            continue
        slide = presentation.slides[number - 1]
        next_id = max(int(p.get("id")) for p in slide._element.iter(qn("p:cNvPr"))) + 1
        for shape in list(slide.shapes):
            if _is_header(shape):
                slide.shapes._spTree.remove(shape._element)
        for index, original in enumerate(reference_elements):
            element = deepcopy(original)
            properties = element.find(".//" + qn("p:cNvPr"))
            properties.set("id", str(next_id + index))
            properties.set("name", HEADER_NAMES[index])
            for node in element.iter():
                for attribute in (qn("r:embed"), qn("r:link"), qn("r:id")):
                    old_id = node.get(attribute)
                    if old_id is None:
                        continue
                    relation = reference.part.rels[old_id]
                    target = relation.target_ref if relation.is_external else relation.target_part
                    node.set(attribute, slide.part.relate_to(
                        target, relation.reltype, is_external=relation.is_external,
                    ))
            if index == 4:
                _set_title(element, titles[number], number)
            if index == 6:
                run_properties = element.find(".//" + qn("a:rPr"))
                hyperlink = run_properties.find(qn("a:hlinkClick"))
                run_properties.remove(hyperlink)
                run_properties.set("u", "sng")
                properties.insert(0, hyperlink)
            slide.shapes._spTree.insert_element_before(element, "p:extLst")
    validate_uniform_headers(presentation)
    return {
        "version": HEADER_DESIGN_VERSION,
        "reference_slide": 7,
        "slide_count": 129,
        "excluded_slides": [1, 131],
    }


def validate_uniform_headers(presentation) -> None:
    for number in range(2, len(presentation.slides)):
        slide = presentation.slides[number - 1]
        shapes = [s for s in slide.shapes if _is_header(s)]
        texts = [_text(s) for s in shapes]
        if len(shapes) != 7 or CONSULTATION not in texts or BIT_LABEL not in texts:
            raise RuntimeError(f"La barra común no se conservó en la diapositiva {number}.")
        if any(s.name == "ANUARIO_BARRA_URL" for s in shapes):
            url = next(s for s in shapes if s.name == "ANUARIO_BARRA_URL")
            hyperlink = url._element.find(".//" + qn("p:cNvPr") + "/" + qn("a:hlinkClick"))
            if hyperlink is None or slide.part.rels[hyperlink.get(qn("r:id"))].target_ref != BIT_URL:
                raise RuntimeError(f"Falta el enlace BIT CRT en la diapositiva {number}.")
