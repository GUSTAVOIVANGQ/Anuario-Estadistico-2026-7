"""Restaura los anexos del anuario de referencia en la edición 2026.

Las tablas y los textos siguen siendo objetos editables de PowerPoint. Los
indicadores estatales que esta edición sí recalcula sustituyen sus columnas
históricas, y la leyenda al pie identifica las distintas fechas de corte.
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import pymupdf
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


INK = RGBColor(31, 48, 49)
TEAL = RGBColor(31, 69, 68)
MID = RGBColor(78, 130, 123)
PALE = RGBColor(228, 240, 237)
MUTED = RGBColor(75, 88, 88)
WHITE = RGBColor(255, 255, 255)


def _clear(slide) -> None:
    for shape in list(slide.shapes):
        if shape.name.startswith("ANUARIO_HEADER_") or shape.name == "ANUARIO_PAGE_NUMBER":
            continue
        slide.shapes._spTree.remove(shape._element)


def _text(slide, name: str, value: str, x: float, y: float, w: float,
          h: float, *, size: float = 7.0, bold: bool = False,
          color: RGBColor = INK, align=PP_ALIGN.LEFT, url: str | None = None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    run = p.add_run()
    run.text = value
    run.font.name = "Noto Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if url:
        run.hyperlink.address = url
    return shape


def _reference_pdf(project_root: Path) -> Path:
    source = project_root / "anuarioestadistico2024vf_0.pdf"
    if not source.is_file():
        source = project_root / "assets" / "reference" / "anuario_2024_fuente" / "anuarioestadistico2024vf_0.pdf"
    if not source.is_file():
        raise FileNotFoundError(f"Falta el anuario 2024: {source}")
    return source


def _figure_rows(project_root: Path, figure: str) -> dict[str, dict[str, str]]:
    metadata = json.loads((project_root / "assets" / "presentation" / "narrativas_figuras_2026.json").read_text(encoding="utf-8"))
    source = project_root / metadata["reference_run"] / "datos_usados" / f"{figure}_datos_usados.csv"
    with source.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    key = "ENTIDAD" if "ENTIDAD" in rows[0] else "entidad"
    return {re.sub(r"\s+", " ", row[key]).strip().casefold(): row for row in rows}


def _updated_tables(project_root: Path, page: int, rows: list[list[str]]) -> None:
    # Los valores nacionales no aparecen en los archivos estatales; se dejan
    # sin estimar en lugar de mezclar una tasa nacional histórica con estados
    # recalculados.
    if page == 118:
        columns = [(4, "b_6", "valor"), (5, "b_7", "valor"),
                   (6, "b_13", "valor"), (7, "b_14", "valor"),
                   (8, "b_21", "penetracion_grafica"),
                   (9, "b_22", "penetracion_grafica")]
    elif page == 119:
        columns = [(14, "c_7", "valor"), (15, "c_13", "valor")]
    else:
        return
    lookup = {figure: _figure_rows(project_root, figure) for _, figure, _ in columns}
    for row in rows:
        entity = re.sub(r"\s+", " ", str(row[0] or "")).strip().casefold()
        if entity == "nacional":
            for column, _, _ in columns:
                row[column] = "—"
            continue
        for column, figure, field in columns:
            value = lookup[figure].get(entity)
            if value is None:
                raise ValueError(f"{page}: falta {entity} en {figure}")
            row[column] = str(round(float(value[field])))


def _table_headers(page: int):
    if page == 117:
        return [
            [(0, 0, "Entidad"), (1, 6, "Indicadores sociodemográficos"),
             (7, 8, "Empleo en telecomunicaciones y radiodifusión"),
             (9, 13, "Equipamiento de TIC en hogares")],
            [(0, 0, ""), (1, 3, "Hogares"), (4, 6, "Habitantes"),
             (7, 7, "Telecom."), (8, 8, "Radiodif."),
             (9, 9, "Radio"), (10, 10, "TV analóg."),
             (11, 11, "TV digital"), (12, 12, "Cómputo"), (13, 13, "Móvil")],
            [(0, 0, ""), (1, 1, "Total"), (2, 2, "Rural"), (3, 3, "Urbano"),
             (4, 4, "Total"), (5, 5, "Rural"), (6, 6, "Urbano")],
        ]
    if page == 118:
        return [
            [(0, 0, "Entidad"), (1, 3, "Personas usuarias de TIC · histórico"),
             (4, 10, "Servicios fijos"), (11, 14, "Internet residencial por tecnología · histórico")],
            [(1, 1, "Internet"), (2, 2, "Computadora"), (3, 3, "Smartphone"),
             (4, 5, "Telefonía fija · actualizado"),
             (6, 7, "Internet fijo · actualizado"),
             (8, 9, "TV restringida · actualizado"),
             (10, 10, "Telefonía pública · histórico"),
             (11, 11, "Cable"), (12, 12, "DSL"), (13, 13, "Fibra"), (14, 14, "Otras")],
            [(4, 4, "Res./100 hogares"), (5, 5, "No res./100 UE"),
             (6, 6, "Res./100 hogares"), (7, 7, "No res./100 UE"),
             (8, 8, "Res./100 hogares"), (9, 9, "No res./100 UE"),
             (10, 10, "Líneas/1,000 hab.")],
        ]
    if page == 119:
        return [
            [(0, 0, "Entidad"), (1, 13, "Telecomunicaciones fijas · datos históricos"),
             (14, 15, "Telecomunicaciones móviles · actualizado")],
            [(1, 4, "Internet no residencial por tecnología"),
             (5, 7, "TV residencial por tecnología"),
             (8, 10, "TV no residencial por tecnología"),
             (11, 13, "Servicios fijos en hogares"),
             (14, 14, "Telefonía móvil"), (15, 15, "Internet móvil")],
            [(1, 1, "Cable"), (2, 2, "DSL"), (3, 3, "Fibra"), (4, 4, "Otras"),
             (5, 5, "Cable"), (6, 6, "Satélite"), (7, 7, "Otras"),
             (8, 8, "Cable"), (9, 9, "Satélite"), (10, 10, "Otras"),
             (11, 11, "Uno"), (12, 12, "Dos"), (13, 13, "Tres"),
             (14, 14, "Líneas/100 hab."), (15, 15, "Líneas/100 hab.")],
        ]
    return [
        [(0, 0, "Entidad"), (1, 14, "Concesiones de radiodifusión · cuadro histórico del Anuario 2024")],
        [(1, 5, "Radio AM"), (6, 10, "Radio FM"), (11, 14, "Televisión digital terrestre")],
        [(1, 1, "Comerc."), (2, 2, "Públicas"), (3, 3, "Sociales"),
         (4, 4, "Comunit."), (5, 5, "Indígenas"),
         (6, 6, "Comerc."), (7, 7, "Públicas"), (8, 8, "Sociales"),
         (9, 9, "Comunit."), (10, 10, "Indígenas"),
         (11, 11, "Comerc."), (12, 12, "Públicas"),
         (13, 13, "Sociales"), (14, 14, "Comunit.")],
    ]


def _source_note(page: int) -> str:
    return {
        117: "Fuente: Anuario Estadístico 2024, Anexo I. Cuadro histórico; sus cortes se indican en los encabezados de origen.",
        118: "Fuente: BIT, dic. 2024, con denominadores ENDUTIH 2025 y DENUE en columnas de servicios (figs. B.6, B.7, B.13, B.14, B.21 y B.22). Demás columnas: Anuario 2024. —: tasa nacional no recalculada.",
        119: "Fuente: BIT, dic. 2024, y denominador de las figuras C.7 y C.13 en las dos columnas móviles. Demás columnas: Anuario 2024. —: tasa nacional no recalculada.",
        120: "Fuente: Anuario Estadístico 2024, Anexo I. Cuadro histórico de concesiones; no se presenta como dato de 2026.",
    }[page]


def _annex_table(slide, source_page, project_root: Path, page: int) -> None:
    _clear(slide)
    raw = source_page.find_tables().tables[0].extract()
    header_rows = 4 if page == 117 else 3
    data = [[str(cell or "").replace("\n", " ") for cell in row] for row in raw[header_rows:]]
    if page == 117:
        # El extractor PDF reconoce una columna final vacía, ajena al cuadro.
        data = [row[:14] for row in data]
    _updated_tables(project_root, page, data)
    headers = _table_headers(page)
    ncols = len(data[0])
    y, h = 0.87, 4.03
    x, w = 0.34, 9.32
    table_shape = slide.shapes.add_table(len(data) + 3, ncols,
                                         Inches(x), Inches(y), Inches(w), Inches(h))
    table_shape.name = f"ANUARIO_ANEXO_I_COMPLETO_{page}"
    table = table_shape.table
    table.first_row = False
    table.horz_banding = False
    first = 1.10 if page != 119 else 1.04
    table.columns[0].width = Inches(first)
    for col in list(table.columns)[1:]:
        col.width = Inches((w - first) / (ncols - 1))
    heights = [0.23, 0.25, 0.27]
    data_height = (h - sum(heights)) / len(data)
    for index, row in enumerate(table.rows):
        row.height = Inches(heights[index] if index < 3 else data_height)
    for ri in range(len(data) + 3):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.margin_left = Inches(0.008)
            cell.margin_right = Inches(0.008)
            cell.margin_top = cell.margin_bottom = 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                TEAL if ri == 0 else MID if ri == 1 else PALE if ri == 2
                else WHITE if ri % 2 else RGBColor(244, 248, 247)
            )
            cell.text_frame.word_wrap = ri < 3
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            if ri >= 3:
                cell.text = data[ri - 3][ci]
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                    p.space_before = p.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = "Noto Sans"
                        run.font.size = Pt(4.95 if ci else 5.05)
                        run.font.bold = ri == 3
                        run.font.color.rgb = INK
    for ri, groups in enumerate(headers):
        for start, end, label in groups:
            cell = table.cell(ri, start)
            if end > start:
                cell.merge(table.cell(ri, end))
                cell = table.cell(ri, start)
            cell.text = label
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                p.space_before = p.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Noto Sans"
                    run.font.size = Pt(5.35 if ri != 2 else 4.95)
                    run.font.bold = True
                    run.font.color.rgb = WHITE if ri < 2 else TEAL
    _text(slide, f"ANUARIO_ANEXO_I_FUENTE_{page}", _source_note(page),
          0.37, 4.96, 9.22, 0.22, size=5.75, color=MUTED)


def _update_reference_text(page: int, text: str) -> str:
    text = (text.replace("\ufb01", "fi").replace("\ufb02", "fl")
            .replace("\x08", "").replace("\t", " "))
    if page == 121:
        text = text.replace("https://www.inegi.org.mx/app/mapa/denue/default.aspx",
                            "https://www.inegi.org.mx/app/mapa/denue/")
    if page == 122:
        text = text.replace("dutih/2022/", "endutih/2025/")
        text = text.replace("enigh/nc/2022/", "enigh/nc/2024/")
    if page == 124:
        text = text.replace("precios constantes de 2013", "precios constantes de 2018")
    if page == 126:
        text = text.replace("ENDUTIH 2023", "ENDUTIH 2025")
        text = text.replace("ENDUTIH 2022, ", "")
        text = text.replace("al IFT", "a la CRT")
    if page == 127:
        text = text.replace("ENIGH 2022", "ENIGH 2024")
    if page == 128:
        text = text.replace("ENDUTIH 2023", "ENDUTIH 2025")
        text = text.replace("2023, PARTICULARIDADES", "2025, PARTICULARIDADES")
        text = text.replace("programas/dutih/2022/", "programas/endutih/2025/")
    return text


def _reference_line(slide, line: dict, spans: list[dict], page: int, serial: int) -> None:
    x0, y0, x1, y1 = line["bbox"]
    sx = 9.16 / 1504
    sy = 4.38 / 750
    left = 0.42 + (x0 - 46) * sx
    top = 0.67 + (y0 - 85) * sy
    width = min(9.70 - left, max(0.12, (x1 - x0) * sx + 0.13))
    height = max(0.10, (y1 - y0) * sy + 0.045)
    if width <= 0:
        return
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.name = f"ANUARIO_ANEXO_P{page}_L{serial:03d}"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.space_before = p.space_after = Pt(0)
    for span in spans:
        run = p.add_run()
        run.text = _update_reference_text(page, span["text"])
        run.font.name = "Noto Sans"
        factor = 0.37 if page in (123, 127, 128) else 0.40
        run.font.size = Pt(round(span["size"] * factor, 2))
        run.font.bold = bool(span["flags"] & 16)
        run.font.italic = bool(span["flags"] & 2)
        source_color = span["color"]
        run.font.color.rgb = (
            INK if source_color == 0 else
            TEAL if source_color >= 0xDDDDDD else
            RGBColor((source_color >> 16) & 255,
                     (source_color >> 8) & 255, source_color & 255)
        )
        if "https://" in run.text:
            match = re.search(r"https://[^\s)]+", run.text)
            if match:
                run.hyperlink.address = match.group(0).rstrip(".,")


def _reference_body(slide, source_page, page: int) -> None:
    _clear(slide)
    if page in (125, 127):
        _reference_table_guides(slide, source_page)
    serial = 0
    for block in source_page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            x0, y0, _, y1 = line["bbox"]
            if y0 < 85 or y0 >= 836 or y1 > 845 or x0 < 40:
                continue
            spans = [span for span in line["spans"] if span["text"].strip()]
            if not spans:
                continue
            serial += 1
            _reference_line(slide, line, spans, page, serial)


def _reference_table_guides(slide, source_page) -> None:
    """Devuelve la retícula de los cuadros de universos e INPC."""
    sx = 9.16 / 1504
    sy = 4.38 / 750

    def mapped(x: float, y: float) -> tuple[float, float]:
        return 0.42 + (x - 46) * sx, 0.67 + (y - 85) * sy

    for index, table in enumerate(source_page.find_tables().tables):
        x0, y0, x1, y1 = table.bbox
        lx, ty = mapped(x0, y0)
        rx, by = mapped(x1, y1)
        border = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(lx), Inches(ty),
            Inches(rx - lx), Inches(by - ty),
        )
        border.name = f"ANUARIO_TABLA_REFERENCIA_{index}"
        border.fill.background()
        border.line.color.rgb = RGBColor(211, 222, 220)
        border.line.width = Pt(0.45)
        ys = sorted({round(v, 2) for cell in table.cells if cell for v in (cell[1], cell[3])})
        for y in ys[1:-1]:
            _, py = mapped(x0, y)
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                              Inches(lx), Inches(py),
                                              Inches(rx), Inches(py))
            line.line.color.rgb = RGBColor(221, 230, 228)
            line.line.width = Pt(0.35)


def _formula_table(slide, name: str, x: float, rows: list[tuple[str, str]]) -> None:
    shape = slide.shapes.add_table(len(rows) + 1, 2,
                                   Inches(x), Inches(1.22), Inches(4.50), Inches(3.69))
    shape.name = name
    table = shape.table
    table.columns[0].width = Inches(1.78)
    table.columns[1].width = Inches(2.72)
    table.rows[0].height = Inches(0.29)
    for row in list(table.rows)[1:]:
        row.height = Inches(3.40 / len(rows))
    for ri, values in enumerate([("Indicador", "Método de cálculo"), *rows]):
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.text = value
            cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if ri == 0 else (WHITE if ri % 2 else RGBColor(242, 247, 246))
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                p.space_before = p.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Noto Sans"
                    run.font.size = Pt(7.0 if ri == 0 else (6.55 if ci == 0 else 6.7))
                    run.font.bold = ri == 0
                    run.font.color.rgb = WHITE if ri == 0 else INK


def _methods_page(slide) -> None:
    _clear(slide)
    _text(slide, "ANUARIO_ANEXO_IV_TITLE", "ANEXO IV. MÉTODOS DE CÁLCULO DE INDICADORES",
          0.42, 0.73, 9.1, 0.29, size=12.0, bold=True, color=TEAL)
    _text(slide, "ANUARIO_ANEXO_IV_INTRO",
          "Para obtener los indicadores mostrados en el documento se utilizaron los siguientes métodos de cálculo:",
          0.42, 1.07, 9.1, 0.17, size=7.5)
    _formula_table(slide, "ANUARIO_ANEXO_IV_TABLA_IZQUIERDA", 0.42, [
        ("Índice de Herfindahl-Hirschman (IHH)", "Σ (participación de mercado de cada grupo económico × 100)²"),
        ("Líneas o accesos residenciales por cada 100 hogares", "Accesos o líneas residenciales ÷ hogares de ENDUTIH del corte × 100"),
        ("Líneas o accesos no residenciales por cada 100 UE", "Accesos o líneas no residenciales ÷ unidades económicas de DENUE × 100"),
        ("Líneas móviles de telefonía por cada 100 habitantes", "Líneas del servicio móvil de telefonía ÷ población compatible con el corte × 100"),
        ("Líneas de Internet móvil por cada 100 habitantes", "Líneas del servicio móvil de Internet ÷ población compatible con el corte × 100"),
    ])
    _formula_table(slide, "ANUARIO_ANEXO_IV_TABLA_DERECHA", 5.08, [
        ("Minutos de uso promedio mensual (MOU)", "Σ (minutos mensuales ÷ líneas del servicio de telefonía de cada mes) ÷ 12"),
        ("Tráfico mensual promedio de Internet móvil", "Σ (GB mensuales ÷ líneas de Internet móvil de cada mes) ÷ 12"),
        ("Mensajes de texto promedio (SMS)", "Σ (SMS enviados por mes ÷ líneas móviles de cada mes) ÷ 12"),
        ("Margen de utilidad", "Ingresos brutos − egresos operativos"),
    ])
    _text(slide, "ANUARIO_ANEXO_IV_NOTA",
          "Nota: hogares y población corresponden a la ENDUTIH del corte de cada figura; las UE provienen del DENUE. Las figuras identifican fuente y periodo.",
          0.43, 4.98, 9.10, 0.18, size=6.2, color=MUTED)


def _back_cover(slide) -> None:
    shape = next((s for s in slide.shapes if s.name == "ANUARIO_BACK_CONTACT"), None)
    if shape is None:
        raise ValueError("Falta el campo de contacto de la contraportada")
    shape.top = Inches(3.69)
    shape.height = Inches(0.78)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    lines = [
        "Comisión Reguladora de Telecomunicaciones · www.gob.mx/crt",
        "Av. Insurgentes Sur 1143, Col. Nochebuena, Benito Juárez, C.P. 03720, Ciudad de México · Tel. 55 5015 4000",
        "Banco de Información de Telecomunicaciones: https://bit.crt.gob.mx/BitWebApp/",
    ]
    for index, value in enumerate(lines):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = p.space_after = Pt(0)
        run = p.add_run()
        run.text = value
        run.font.name = "Noto Sans"
        run.font.size = Pt(7.2 if index == 0 else 6.65)
        run.font.bold = index == 0
        run.font.color.rgb = RGBColor(239, 247, 243)


def apply_annex_pages(project_root: Path, presentation) -> None:
    """Reemplaza únicamente las diapositivas 117–131."""
    with pymupdf.open(_reference_pdf(project_root)) as reference:
        for page in range(117, 121):
            _annex_table(presentation.slides[page - 1], reference[page - 1], project_root, page)
        for page in range(121, 130):
            _reference_body(presentation.slides[page - 1], reference[page - 1], page)
    _methods_page(presentation.slides[129])
    _back_cover(presentation.slides[130])
