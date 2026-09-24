from __future__ import annotations

import base64
from io import BytesIO
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image

from .figure_outputs import companion_paths, extract_svg_text, validate_editable_svg
from .narratives import load_narrative_audit, write_narrative_reports
from .presentation_svg import SvgEmbeddingRequest, embed_svg_images


DEFAULT_INSET_INCHES = 0.08
TRANSPARENT_BOOTSTRAP_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


@dataclass(frozen=True)
class AssemblyResult:
    output_path: Path
    report_path: Path
    inserted: int
    missing: tuple[str, ...]
    errors: tuple[str, ...]
    slide_count: int


def _load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as stream:
        return json.load(stream)


def _resolve_presentation_config(project_root: Path, config: dict[str, Any]) -> tuple[Path, Path]:
    presentation = config.get("presentation", {})
    template_rel = presentation.get(
        "template", "assets/presentation/anuario_estadistico_2026_automatizable.pptx"
    )
    manifest_rel = presentation.get(
        "manifest", "assets/presentation/anuario_estadistico_2026_manifest.json"
    )
    return project_root / template_rel, project_root / manifest_rel


def _shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def _insert_editorial_text(project_root: Path, presentation, manifest: dict[str, Any]) -> dict[str, Any]:
    """Fill named editorial boxes when reviewed copy is available."""

    registry_path = project_root / "assets" / "presentation" / "textos_editoriales_2026.json"
    slots = manifest.get("editorial_slots", [])
    registry = _load_json(registry_path) if registry_path.is_file() else {}
    values = registry.get("slots", {})
    source_notes = registry.get("source_notes", {})
    inserted: list[str] = []
    pending: list[str] = []
    for slot in slots:
        key = str(slot["key"])
        value = values.get(key)
        if not isinstance(value, str) or not value.strip():
            pending.append(key)
            continue
        shape = _shape_by_name(presentation.slides[int(slot["slide_number"]) - 1], slot["shape_name"])
        if shape is None or not shape.has_text_frame:
            raise ValueError(f"Falta el campo editorial {key}: {slot['shape_name']}")
        token = str(slot["token"])
        if slot.get("kind") == "tabla":
            rows = [line.split("\t") for line in value.strip().splitlines() if line.strip()]
            if not rows or len({len(row) for row in rows}) != 1:
                raise ValueError(f"La tabla editorial {key} no es rectangular")
            shape.text = shape.text.replace(token, "")
            inset = Inches(0.12)
            top_offset = Inches(0.47)
            source_note = str(source_notes.get(key) or "")
            note_space = Inches(0.23) if source_note else 0
            table_shape = presentation.slides[int(slot["slide_number"]) - 1].shapes.add_table(
                len(rows), len(rows[0]),
                shape.left + inset, shape.top + top_offset,
                shape.width - 2 * inset, shape.height - top_offset - inset - note_space,
            )
            table_shape.name = "ANUARIO_TABLE_" + key
            table = table_shape.table
            table.first_row = False
            table.horz_banding = False
            first_width = int((shape.width - 2 * inset) * (0.24 if len(rows[0]) > 3 else 0.30))
            table.columns[0].width = first_width
            for column in list(table.columns)[1:]:
                column.width = int((shape.width - 2 * inset - first_width) / (len(rows[0]) - 1))
            font_size = 5.4 if len(rows) > 25 else (8.3 if len(rows) > 10 else 9.0)
            for row_index, row_values in enumerate(rows):
                for column_index, cell_text in enumerate(row_values):
                    cell = table.cell(row_index, column_index)
                    cell.text = cell_text
                    cell.margin_left = Inches(0.015)
                    cell.margin_right = Inches(0.015)
                    cell.margin_top = 0
                    cell.margin_bottom = 0
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = (
                        RGBColor(32, 61, 62) if row_index == 0 else
                        RGBColor(245, 247, 247) if row_index % 2 == 0 else RGBColor(255, 255, 255)
                    )
                    cell.text_frame.word_wrap = False
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.space_before = Pt(0)
                        paragraph.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.name = "Noto Sans"
                            run.font.size = Pt(font_size)
                            run.font.bold = row_index == 0
                            run.font.color.rgb = (
                                RGBColor(255, 255, 255) if row_index == 0 else RGBColor(32, 54, 52)
                            )
            if source_note:
                note = presentation.slides[int(slot["slide_number"]) - 1].shapes.add_textbox(
                    shape.left + inset,
                    shape.top + shape.height - Inches(0.23),
                    shape.width - 2 * inset,
                    Inches(0.20),
                )
                note.name = "ANUARIO_SOURCE_" + key
                note.text = source_note
                note.text_frame.margin_left = 0
                note.text_frame.margin_top = 0
                for paragraph in note.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Noto Sans"
                        run.font.size = Pt(6.2)
                        run.font.color.rgb = RGBColor(67, 67, 67)
            inserted.append(key)
            continue
        replaced = False
        area = (shape.width / 914400) * (shape.height / 914400)
        fitted_size = max(6.0, min(8.5, math.sqrt(area * 6200 / len(value))))
        if key.startswith("ANEXO_II_"):
            # The source definitions fill ten dense columns. Keep the complete
            # wording within the fixed 2024 page sequence and avoid overflow.
            fitted_size = 5.6
            shape.text_frame.margin_top = Inches(0.02)
            shape.text_frame.margin_bottom = Inches(0.02)
            shape.text_frame.margin_left = Inches(0.02)
            shape.text_frame.margin_right = Inches(0.02)
            shape.text_frame.word_wrap = True
        for paragraph in shape.text_frame.paragraphs:
            if key.startswith("ANEXO_II_"):
                paragraph.space_before = Pt(0)
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
            for run in paragraph.runs:
                if token in run.text:
                    run.text = run.text.replace(token, value.strip())
                    current_size = run.font.size.pt if run.font.size else fitted_size
                    run.font.size = Pt(min(current_size, fitted_size))
                    replaced = True
        if not replaced:
            if token not in shape.text:
                raise ValueError(f"Falta el token editorial {token} en {slot['shape_name']}")
            shape.text = shape.text.replace(token, value.strip())
        inserted.append(key)
    return {"registry": str(registry_path), "inserted": inserted, "pending": pending}


def _narrative_font_size(text: str) -> float:
    """Return a conservative body size for the fixed narrative card.

    The cards are intentionally the same size on every slide.  The source PDF
    contains anything from a one-line observation to multi-paragraph notes, so
    a small deterministic scale is safer than relying on viewer-specific
    PowerPoint autofit behaviour.
    """

    length = len(text)
    if length <= 260:
        return 8.8
    if length <= 430:
        return 8.2
    if length <= 650:
        return 7.6
    if length <= 900:
        return 7.0
    if length <= 1_180:
        return 6.5
    return 6.0


def _set_narrative_shape_text(
    shape, narrative: str, *, context_line: str | None = None
) -> dict[str, Any]:
    """Replace only the template filler with reviewed native PPTX prose."""

    if not getattr(shape, "has_text_frame", False):
        raise ValueError(f"El objeto {shape.name!r} no admite texto.")

    original_lines = [line.strip() for line in shape.text.splitlines() if line.strip()]
    label = original_lines[0] if original_lines else shape.name
    title = original_lines[1] if len(original_lines) > 1 else ""
    if title.startswith("{{TEXTO:"):
        title = ""
    body_size = _narrative_font_size(narrative)

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP

    label_paragraph = frame.paragraphs[0]
    label_paragraph.text = label
    label_paragraph.space_before = Pt(0)
    label_paragraph.space_after = Pt(2)
    for run in label_paragraph.runs:
        run.font.name = "Noto Sans"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(32, 54, 52)

    if title:
        title_paragraph = frame.add_paragraph()
        title_paragraph.text = title
        title_paragraph.space_before = Pt(0)
        title_paragraph.space_after = Pt(6)
        for run in title_paragraph.runs:
            run.font.name = "Noto Sans"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(32, 54, 52)

    if context_line:
        context = frame.add_paragraph()
        context.text = context_line
        context.space_after = Pt(5)
        for run in context.runs:
            run.font.name = "Noto Sans"
            run.font.size = Pt(7)
            run.font.bold = True
            run.font.color.rgb = RGBColor(74, 125, 117)

    for paragraph_text in [part.strip() for part in narrative.split("\n\n") if part.strip()]:
        paragraph = frame.add_paragraph()
        paragraph.text = paragraph_text
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(4)
        paragraph.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Noto Sans"
            run.font.size = Pt(body_size)
            run.font.bold = False
            run.font.color.rgb = RGBColor(67, 67, 67)

    return {
        "label": label,
        "title": title,
        "body_font_size_pt": body_size,
        "character_count": len(narrative),
        "paragraph_count": len([x for x in narrative.split("\n\n") if x.strip()]),
    }


def _contain_box(
    svg_path: Path,
    left: int,
    top: int,
    width: int,
    height: int,
    *,
    inset_inches: float,
) -> tuple[int, int, int, int]:
    inset = int(Inches(inset_inches))
    box_left = left + inset
    box_top = top + inset
    box_width = max(1, width - 2 * inset)
    box_height = max(1, height - 2 * inset)

    root = ElementTree.parse(svg_path).getroot()
    view_box = (root.get("viewBox") or "").replace(",", " ").split()
    if len(view_box) == 4:
        svg_width = float(view_box[2])
        svg_height = float(view_box[3])
    else:
        def dimension(value: str | None) -> float:
            if not value:
                raise ValueError(f"El SVG no declara dimensiones: {svg_path}")
            numeric = "".join(character for character in value if character in "0123456789+-.eE")
            return float(numeric)

        svg_width = dimension(root.get("width"))
        svg_height = dimension(root.get("height"))
    if svg_width <= 0 or svg_height <= 0:
        raise ValueError(f"Dimensiones inválidas en {svg_path}")

    image_ratio = svg_width / svg_height
    box_ratio = box_width / box_height
    if image_ratio >= box_ratio:
        target_width = box_width
        target_height = max(1, int(round(target_width / image_ratio)))
    else:
        target_height = box_height
        target_width = max(1, int(round(target_height * image_ratio)))

    target_left = box_left + (box_width - target_width) // 2
    target_top = box_top + (box_height - target_height) // 2
    return target_left, target_top, target_width, target_height


def _add_searchable_text_layer(
    slide,
    bounds: tuple[int, int, int, int],
    figure_id: str,
    texts: tuple[str, ...],
):
    left, top, width, height = bounds
    marker = f"Texto seleccionable de la figura {figure_id}"
    layer = slide.shapes.add_textbox(left, top, width, height)
    layer.name = "ANUARIO_TEXT_LAYER_" + figure_id.replace(".", "_")
    layer.fill.background()
    layer.line.fill.background()
    layer._element.nvSpPr.cNvPr.set(
        "descr",
        f"Capa accesible con el texto del SVG de la figura {figure_id}",
    )

    text_frame = layer.text_frame
    text_frame.clear()
    # Un párrafo por elemento evita que PowerPoint introduzca espacios dentro
    # de las palabras al generar el mapa de texto del PDF.
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.text = "\n".join((marker, *texts))
    for paragraph in text_frame.paragraphs:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Noto Sans"
            run.font.size = Pt(1)
            run.font.color.rgb = RGBColor(255, 255, 255)
    return layer, marker


def _insert_picture_after_placeholder(
    slide,
    placeholder,
    svg_path: Path,
    figure_id: str,
    *,
    searchable_texts: tuple[str, ...] = (),
) -> tuple[str | None, int]:
    left, top, width, height = _contain_box(
        svg_path,
        placeholder.left,
        placeholder.top,
        placeholder.width,
        placeholder.height,
        inset_inches=DEFAULT_INSET_INCHES,
    )

    if getattr(placeholder, "has_text_frame", False):
        placeholder.text_frame.clear()

    sp_tree = slide.shapes._spTree
    placeholder_index = sp_tree.index(placeholder._element)
    text_layer = None
    marker = None
    if searchable_texts:
        text_layer, marker = _add_searchable_text_layer(
            slide,
            (left, top, width, height),
            figure_id,
            searchable_texts,
        )
    # ``python-pptx`` todavía no admite SVG. Se crea la geometría con un PNG
    # transparente de 1x1 y presentation_svg sustituye después esa relación y
    # elimina el medio temporal del paquete final.
    picture = slide.shapes.add_picture(
        BytesIO(TRANSPARENT_BOOTSTRAP_PNG),
        left,
        top,
        width,
        height,
    )
    picture.name = "ANUARIO_IMAGE_" + figure_id.replace(".", "_")
    picture._element.nvPicPr.cNvPr.set(
        "descr",
        f"Figura {figure_id}. SVG vectorial sin respaldo raster",
    )

    picture_element = picture._element
    sp_tree.remove(picture_element)
    if text_layer is not None:
        text_element = text_layer._element
        sp_tree.remove(text_element)
        sp_tree.insert(placeholder_index + 1, text_element)
        sp_tree.insert(placeholder_index + 2, picture_element)
    else:
        sp_tree.insert(placeholder_index + 1, picture_element)
    return marker, len(searchable_texts)


def _validate_native_text_layers(
    presentation: Presentation,
    inserted_rows: list[dict[str, Any]],
) -> None:
    """Confirma que el PPTX conserva cada texto SVG como texto nativo copiable."""
    for row in inserted_rows:
        slide = presentation.slides[int(row["slide_number"]) - 1]
        figure_id = str(row["figure_id"])
        layer_name = "ANUARIO_TEXT_LAYER_" + figure_id.replace(".", "_")
        layer = _shape_by_name(slide, layer_name)
        if layer is None or not getattr(layer, "has_text_frame", False):
            raise RuntimeError(f"El PPTX perdió la capa de texto de {figure_id}.")
        native_text = layer.text
        expected = [str(row["searchable_text_marker"]), *row["svg_texts"]]
        missing = [text for text in expected if text not in native_text]
        if missing:
            raise RuntimeError(
                f"El PPTX perdió {len(missing)} textos copiables de {figure_id}."
            )


def _validate_narrative_layers(
    presentation: Presentation,
    inserted_rows: list[dict[str, Any]],
) -> None:
    """Confirm that each reviewed paragraph remains native, selectable text."""

    for row in inserted_rows:
        narrative_shape_name = row.get("narrative_shape_name")
        narrative_text = row.get("narrative_text")
        if not narrative_shape_name or not narrative_text:
            continue
        slide = presentation.slides[int(row["slide_number"]) - 1]
        shape = _shape_by_name(slide, str(narrative_shape_name))
        if shape is None or not getattr(shape, "has_text_frame", False):
            raise RuntimeError(
                f"El PPTX perdió la narrativa nativa de {row['figure_id']}."
            )
        expected = re.sub(r"\s+", " ", str(narrative_text)).strip()
        actual = re.sub(r"\s+", " ", shape.text).strip()
        if expected not in actual:
            raise RuntimeError(
                f"El PPTX no conserva íntegro el párrafo copiable de {row['figure_id']}."
            )


def assemble_from_template(
    project_root: Path,
    output_path: Path,
    *,
    strict: bool = False,
    run_dir: Path | None = None,
) -> AssemblyResult:
    """Inserta las figuras generadas en la plantilla automatizable del Anuario 2026.

    La posición de cada figura se obtiene del manifest y el objeto de PowerPoint se localiza
    por su nombre estable (por ejemplo, ``ANUARIO_FIGURE_A_1``). La tarjeta de fondo se
    conserva para mantener el diseño; sólo se borra el token visual y se coloca la imagen
    encima respetando proporción (fit=contain).
    """

    config = _load_json(project_root / "config" / "proyecto.json")
    template_path, manifest_path = _resolve_presentation_config(project_root, config)
    if not template_path.is_file():
        raise FileNotFoundError(f"Falta la plantilla PowerPoint automatizable: {template_path}")
    if not manifest_path.is_file():
        raise FileNotFoundError(f"Falta el manifest de la plantilla: {manifest_path}")

    manifest = _load_json(manifest_path)
    entries = manifest.get("entries", [])
    if not entries:
        raise ValueError(f"El manifest no contiene entradas de figuras: {manifest_path}")

    presentation = Presentation(str(template_path))
    expected_slide_count = int(manifest.get("slide_count") or len(presentation.slides))
    if len(presentation.slides) != expected_slide_count:
        raise ValueError(
            "La plantilla no coincide con el manifest: "
            f"{len(presentation.slides)} diapositivas vs {expected_slide_count}."
        )

    editorial = _insert_editorial_text(project_root, presentation, manifest)

    inserted = 0
    missing: list[str] = []
    errors: list[str] = []
    inserted_rows: list[dict[str, Any]] = []
    svg_requests: list[SvgEmbeddingRequest] = []
    svg_missing: list[str] = []
    searchable_text_layer_count = 0
    narrative_audit = load_narrative_audit(project_root, run_dir=run_dir)
    narrative_inserted: set[str] = set()
    narrative_missing: list[str] = []
    narrative_errors: list[str] = []
    narrative_data_changed: list[str] = []
    historical_rows: list[dict[str, Any]] = []

    for entry in entries:
        figure_id = str(entry["figure_id"])
        slide_number = int(entry["slide_number"])
        shape_name = str(entry["figure_shape_name"])
        narrative_shape_name = str(entry.get("narrative_shape_name") or "")
        image_rel = Path(str(entry["expected_image"]))
        image_path = project_root / image_rel

        svg_path = companion_paths(image_path)["svg"]
        historical_rel = entry.get("historical_image")
        if not svg_path.is_file() and historical_rel:
            historical_path = project_root / str(historical_rel)
            if not historical_path.is_file():
                errors.append(f"{figure_id}: falta el recorte histórico {historical_path}")
                continue
            if slide_number < 1 or slide_number > len(presentation.slides):
                errors.append(f"{figure_id}: diapositiva histórica fuera de rango")
                continue
            slide = presentation.slides[slide_number - 1]
            placeholder = _shape_by_name(slide, shape_name)
            narrative_shape = _shape_by_name(slide, narrative_shape_name)
            record = narrative_audit.records.get(figure_id)
            if placeholder is None or narrative_shape is None or record is None:
                errors.append(f"{figure_id}: faltan marcadores o narrativa históricos")
                continue
            _set_narrative_shape_text(
                narrative_shape,
                record.updated_text,
                context_line="Serie histórica: julio de 2023 a junio de 2024",
            )
            narrative_inserted.add(figure_id)
            placeholder.text_frame.clear()
            with Image.open(historical_path) as bitmap:
                image_ratio = bitmap.width / bitmap.height
            box_ratio = placeholder.width / placeholder.height
            if image_ratio >= box_ratio:
                width = placeholder.width
                height = int(round(width / image_ratio))
            else:
                height = placeholder.height
                width = int(round(height * image_ratio))
            picture = slide.shapes.add_picture(
                str(historical_path),
                placeholder.left + (placeholder.width - width) // 2,
                placeholder.top + (placeholder.height - height) // 2,
                width=width,
                height=height,
            )
            picture.name = "ANUARIO_HISTORICAL_" + figure_id.replace(".", "_")
            historical_rows.append({
                "figure_id": figure_id,
                "slide_number": slide_number,
                "source_pdf_page": entry.get("historical_source_pdf_page"),
                "period": entry.get("historical_period"),
                "image": str(historical_rel),
            })
            continue
        if not svg_path.is_file():
            missing.append(figure_id)
            svg_missing.append(figure_id)
            continue
        if slide_number < 1 or slide_number > len(presentation.slides):
            errors.append(f"{figure_id}: diapositiva fuera de rango ({slide_number})")
            continue

        slide = presentation.slides[slide_number - 1]
        placeholder = _shape_by_name(slide, shape_name)
        if placeholder is None:
            errors.append(
                f"{figure_id}: no se encontró el objeto {shape_name!r} en la diapositiva "
                f"{slide_number}"
            )
            continue

        narrative_record = narrative_audit.records.get(figure_id)
        narrative_meta: dict[str, Any] | None = None
        if narrative_record is None:
            narrative_missing.append(figure_id)
        elif not narrative_shape_name:
            narrative_errors.append(f"{figure_id}: el manifest no declara narrative_shape_name")
        else:
            narrative_shape = _shape_by_name(slide, narrative_shape_name)
            if narrative_shape is None:
                narrative_errors.append(
                    f"{figure_id}: no se encontró el objeto {narrative_shape_name!r} "
                    f"en la diapositiva {slide_number}"
                )
            else:
                try:
                    narrative_meta = _set_narrative_shape_text(
                        narrative_shape,
                        narrative_record.updated_text,
                    )
                    narrative_inserted.add(figure_id)
                    if narrative_record.verification_status == "data_changed":
                        narrative_data_changed.append(figure_id)
                except Exception as exc:  # noqa: BLE001 - reportar narrativa específica
                    narrative_errors.append(
                        f"{figure_id}: narrativa {type(exc).__name__}: {exc}"
                    )

        try:
            svg_inspection = validate_editable_svg(svg_path, require_fully_vector=True)
            searchable_texts = extract_svg_text(svg_path)
        except (OSError, ValueError) as exc:
            svg_missing.append(figure_id)
            errors.append(f"{figure_id}: SVG inválido: {exc}")
            continue

        try:
            searchable_marker, svg_text_count = _insert_picture_after_placeholder(
                slide,
                placeholder,
                svg_path,
                figure_id,
                searchable_texts=searchable_texts,
            )
        except Exception as exc:  # noqa: BLE001 - reportar figura específica
            errors.append(f"{figure_id}: {type(exc).__name__}: {exc}")
            continue

        picture_name = "ANUARIO_IMAGE_" + figure_id.replace(".", "_")
        searchable_text_layer_count += 1
        svg_requests.append(
            SvgEmbeddingRequest(
                figure_id=figure_id,
                slide_number=slide_number,
                picture_name=picture_name,
                svg_path=svg_path,
            )
        )

        inserted += 1
        inserted_rows.append(
            {
                "figure_id": figure_id,
                "slide_number": slide_number,
                "shape_name": shape_name,
                "svg_path": str(svg_path.relative_to(project_root)).replace("\\", "/"),
                "picture_name": picture_name,
                "svg_text_elements": svg_text_count,
                "svg_texts": list(searchable_texts),
                "transparent_background": svg_inspection.transparent_background,
                "fully_vector": svg_inspection.fully_vector,
                "embedded_raster_images": svg_inspection.embedded_images,
                "raster_fallback": False,
                "searchable_text_marker": searchable_marker,
                "narrative_shape_name": narrative_shape_name or None,
                "narrative_text": narrative_record.updated_text if narrative_record else None,
                "narrative_source_pdf_page": (
                    narrative_record.source_pdf_page if narrative_record else None
                ),
                "narrative_update_mode": (
                    narrative_record.update_mode if narrative_record else None
                ),
                "narrative_verification_status": (
                    narrative_record.verification_status if narrative_record else None
                ),
                "narrative_notes": list(narrative_record.notes) if narrative_record else [],
                "narrative_layout": narrative_meta,
            }
        )

    if strict and (
        missing
        or errors
        or svg_missing
        or narrative_missing
        or narrative_errors
        or narrative_data_changed
        or editorial["pending"]
    ):
        details = []
        if missing:
            details.append(f"faltan {len(missing)} figuras: {', '.join(missing)}")
        if errors:
            details.append(f"errores: {'; '.join(errors)}")
        if svg_missing:
            details.append(f"faltan {len(svg_missing)} SVG nativos: {', '.join(svg_missing)}")
        if narrative_missing:
            details.append(
                f"faltan {len(narrative_missing)} narrativas: {', '.join(narrative_missing)}"
            )
        if narrative_errors:
            details.append(f"errores de narrativa: {'; '.join(narrative_errors)}")
        if narrative_data_changed:
            details.append(
                "cambió el archivo de datos revisado para "
                + ", ".join(narrative_data_changed)
            )
        if editorial["pending"]:
            details.append(
                f"faltan {len(editorial['pending'])} campos editoriales: "
                + ", ".join(editorial["pending"])
            )
        raise RuntimeError("No se generó el PPTX en modo estricto; " + " | ".join(details))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()
    presentation.save(str(output_path))
    embeddings = embed_svg_images(output_path, svg_requests)

    # Verificación mínima del paquete recién escrito.
    reopened = Presentation(str(output_path))
    if len(reopened.slides) != len(presentation.slides):
        raise RuntimeError("La verificación del PPTX detectó un número de diapositivas distinto.")
    _validate_native_text_layers(reopened, inserted_rows)
    _validate_narrative_layers(reopened, inserted_rows)

    narrative_report_json = output_path.with_name(output_path.stem + "_narrativas.json")
    narrative_report_csv = output_path.with_name(output_path.stem + "_narrativas.csv")
    write_narrative_reports(
        narrative_audit,
        json_path=narrative_report_json,
        csv_path=narrative_report_csv,
        inserted_figure_ids=narrative_inserted,
    )

    report_path = output_path.with_name(output_path.stem + "_ensamblaje.json")
    report = {
        "project": config.get("project_name", "Anuario Estadístico 2026"),
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "template": str(template_path.relative_to(project_root)).replace("\\", "/"),
        "manifest": str(manifest_path.relative_to(project_root)).replace("\\", "/"),
        "output": str(output_path.relative_to(project_root)).replace("\\", "/")
        if output_path.is_relative_to(project_root)
        else str(output_path),
        "run_dir": str(run_dir.relative_to(project_root)).replace("\\", "/")
        if run_dir and run_dir.is_relative_to(project_root)
        else (str(run_dir) if run_dir else None),
        "slide_count": len(presentation.slides),
        "editorial_inserted_count": len(editorial["inserted"]),
        "editorial_pending_count": len(editorial["pending"]),
        "editorial_pending_keys": editorial["pending"],
        "historical_figure_count": len(historical_rows),
        "historical_figures": historical_rows,
        "manifest_figure_count": len(entries),
        "inserted_count": inserted,
        "embedded_svg_count": len(embeddings),
        "transparent_svg_count": sum(
            1 for row in inserted_rows if row["transparent_background"]
        ),
        "fully_vector_svg_count": sum(
            1 for row in inserted_rows if row["fully_vector"]
        ),
        "embedded_raster_image_count": sum(
            int(row["embedded_raster_images"]) for row in inserted_rows
        ),
        "raster_fallback_count": 0,
        "searchable_text_layer_count": searchable_text_layer_count,
        "narrative_registry": str(narrative_audit.registry_path),
        "narrative_registered_count": len(narrative_audit.records),
        "narrative_inserted_count": len(narrative_inserted),
        "narrative_missing_count": len(narrative_missing),
        "narrative_missing_figures": narrative_missing,
        "narrative_error_count": len(narrative_errors),
        "narrative_errors": narrative_errors,
        "narrative_data_changed_count": len(narrative_data_changed),
        "narrative_data_changed_figures": narrative_data_changed,
        "narrative_verified_count": narrative_audit.verified_count,
        "narrative_update_mode_counts": narrative_audit.mode_counts,
        "narrative_report_json": str(narrative_report_json),
        "narrative_report_csv": str(narrative_report_csv),
        "svg_missing_count": len(svg_missing),
        "svg_missing_figures": svg_missing,
        "missing_count": len(missing),
        "error_count": len(errors),
        "missing_figures": missing,
        "errors": errors,
        "inserted": inserted_rows,
        "fit_policy": "contain",
        "image_policy": (
            "figuras 100% vectoriales en SVG nativo directo, sin PNG/JPG de respaldo "
            "ni imágenes raster incrustadas"
        ),
        "background_policy": (
            "lienzo SVG transparente; se conservan sólo fondos explícitos que forman parte "
            "del diseño de la figura"
        ),
        "text_policy": (
            "texto real <text> en SVG y copia nativa auxiliar en PPTX para búsqueda/copia "
            "en visores de PowerPoint"
        ),
        "narrative_policy": (
            "redacción base del Anuario 2024 conservada; se actualizan registros puntuales "
            "con la figura actual y todo el párrafo se inserta como texto nativo copiable"
        ),
        "placeholder_policy": "conservar tarjeta, limpiar token y superponer figura",
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    return AssemblyResult(
        output_path=output_path,
        report_path=report_path,
        inserted=inserted,
        missing=tuple(missing),
        errors=tuple(errors),
        slide_count=len(presentation.slides),
    )
