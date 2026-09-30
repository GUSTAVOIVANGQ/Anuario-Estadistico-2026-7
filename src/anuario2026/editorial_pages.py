"""Composición de páginas editoriales extensas del Anuario 2026."""

from __future__ import annotations

import json
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


INK = RGBColor(31, 48, 49)
HEADING = RGBColor(72, 73, 121)
MUTED = RGBColor(79, 91, 92)


def _clear_body(slide) -> None:
    for shape in list(slide.shapes):
        if shape.name.startswith("ANUARIO_HEADER_") or shape.name == "ANUARIO_PAGE_NUMBER":
            continue
        slide.shapes._spTree.remove(shape._element)


def _box(slide, name: str, x: float, y: float, w: float, h: float):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(218, 225, 225)
    shape.line.width = Pt(0.8)
    shape.adjustments[0] = 0.08
    return shape


def _text(slide, name: str, value: str, x: float, y: float, w: float, h: float,
          *, size: float, bold: bool = False, color: RGBColor = INK,
          italic_word: str | None = None, hyperlink: str | None = None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    parts = value.split(italic_word) if italic_word else [value]
    for index, part in enumerate(parts):
        if index:
            run = paragraph.add_run()
            run.text = italic_word
            run.font.italic = True
            run.font.name = "Noto Sans"
            run.font.size = Pt(size)
            run.font.color.rgb = color
        if part:
            run = paragraph.add_run()
            run.text = part
            run.font.name = "Noto Sans"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            if hyperlink:
                run.hyperlink.address = hyperlink
    return shape


def _page_81(slide, values: dict[str, str]) -> None:
    _clear_body(slide)
    _text(
        slide, "ANUARIO_F_SOURCE", "Fuente: INEGI, ENDUTIH 2025",
        0.48, 0.73, 5.0, 0.17, size=7.1, color=MUTED,
        hyperlink="https://www.inegi.org.mx/programas/endutih/2025/",
    )
    cards = (
        ("SMARTPHONE", "ACTIVIDADES EN EL TELÉFONO CELULAR INTELIGENTE",
         "F_OVERVIEW_SMARTPHONE", 0.48, 1.02, 2.84, 4.08, 7.6),
        ("INTERNET", "ACTIVIDADES REALIZADAS EN INTERNET",
         "F_OVERVIEW_INTERNET", 3.48, 1.02, 6.02, 1.34, 7.35),
        ("COMPUTADORA", "ACTIVIDADES EN LA COMPUTADORA",
         "F_OVERVIEW_COMPUTADORA", 3.48, 2.48, 6.02, 1.22, 7.55),
        ("REDES", "USO DE LAS REDES SOCIALES",
         "F_OVERVIEW_REDES", 3.48, 3.82, 6.02, 1.28, 7.35),
    )
    for suffix, title, key, x, y, w, h, font_size in cards:
        _box(slide, "ANUARIO_F_CARD_" + suffix, x, y, w, h)
        _text(slide, "ANUARIO_F_TITLE_" + suffix, title,
              x + 0.15, y + 0.12, w - 0.30, 0.32,
              size=9.0, bold=True, color=HEADING)
        _text(slide, "ANUARIO_F_BODY_" + suffix, values[key],
              x + 0.15, y + 0.49, w - 0.30, h - 0.59,
              size=font_size,
              italic_word="Smartphone" if suffix == "SMARTPHONE" else None)


def _page_116(slide, values: dict[str, str]) -> None:
    _clear_body(slide)
    _text(
        slide, "ANUARIO_TOOLS_TITLE",
        "CONOCE LAS HERRAMIENTAS CON INFORMACIÓN POR ENTIDAD FEDERATIVA:",
        0.60, 0.78, 8.86, 0.34, size=12.3, bold=True, color=HEADING,
    )
    tools = (
        ("BIT_ESTATAL", "BIT ESTATAL:", "https://bitestatal.ift.org.mx/",
         values["HERRAMIENTA_1"].split("\n", 1)[1], 1.28),
        ("CALCULADORA", "CALCULADORA DE PROBABILIDADES:",
         "https://calculadoraprob.ift.org.mx/",
         values["HERRAMIENTA_2"].split("\n", 1)[1], 3.02),
    )
    for suffix, title, url, paragraph, y in tools:
        _box(slide, "ANUARIO_TOOL_CARD_" + suffix, 0.60, y, 8.86, 1.58)
        _text(slide, "ANUARIO_TOOL_TITLE_" + suffix, title,
              0.84, y + 0.15, 8.3, 0.22, size=10.5, bold=True, color=HEADING)
        _text(slide, "ANUARIO_TOOL_URL_" + suffix, url,
              0.84, y + 0.44, 8.3, 0.29, size=12.0, bold=True,
              color=RGBColor(41, 107, 140), hyperlink=url)
        _text(slide, "ANUARIO_TOOL_BODY_" + suffix, paragraph,
              0.84, y + 0.85, 8.3, 0.58, size=8.7)
    _text(slide, "ANUARIO_TOOL_CURRENT", values["HERRAMIENTA_3"],
          0.63, 4.80, 8.9, 0.17, size=7.5, color=MUTED,
          hyperlink="https://bit.crt.gob.mx/BitWebApp/")
    _text(slide, "ANUARIO_TOOL_HISTORICAL", values["HERRAMIENTA_4"],
          0.63, 5.02, 8.9, 0.16, size=7.0, color=MUTED)


def apply_editorial_pages(project_root: Path, presentation) -> None:
    registry = project_root / "assets" / "presentation" / "textos_editoriales_2026.json"
    values = json.loads(registry.read_text(encoding="utf-8"))["slots"]
    _page_81(presentation.slides[80], values)
    _page_116(presentation.slides[115], values)
