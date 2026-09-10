"""
Figura A.5 — Inversión Extranjera Directa (IED) en telecomunicaciones
(Versión SOLO DATOS — sin generación de gráfica)

Este script reproduce únicamente la lectura y el cálculo de los datos que
se usarían para construir la Figura A.5, y los imprime en consola.

Datos actualizados al 3er trimestre de 2025.
Período: 2013-2024 (2024 acumulado a junio).

Fuente: Secretaría de Economía – Registro Nacional de Inversiones Extranjeras.
"""

import os
import numpy as np
import openpyxl
from pathlib import Path
import sys

try:
    sys.path.append(str(Path(__file__).resolve().parents[1]))
    from _plot_data_logger import enable_plot_data_logging
    enable_plot_data_logging()
except ImportError:
    pass

# ── 1. Leer datos ─────────────────────────────────────────────────────────────
base = os.path.join(os.path.dirname(__file__), "..", "..", 'datos', 'A.5')

# --- IED total de México (datos actualizados) ---
wb1 = openpyxl.load_workbook(
    os.path.join(base, 'Datos_originales_y_actualizacion__1_.xlsx'),
    data_only=True)
ws1 = wb1['Preliminares y actualización']

total_ied = {}
for r in range(3, ws1.max_row + 1):
    yr = ws1.cell(r, 1).value
    period = ws1.cell(r, 2).value
    val = ws1.cell(r, 4).value  # Columna D = datos actualizados
    if yr and period and val:
        yr = int(yr)
        if 2013 <= yr <= 2024:
            if yr < 2024 and 'diciembre' in str(period):
                total_ied[yr] = float(val)
            elif yr == 2024 and 'junio' in str(period):
                total_ied[yr] = float(val)
wb1.close()

# --- IED en telecomunicaciones (sector 517, datos actualizados) ---
wb2 = openpyxl.load_workbook(
    os.path.join(base, '2025_3T_Flujosportipodeinversion_actu__3_.xlsx'),
    data_only=True, read_only=True)
ws2 = wb2['Por sector']

YEAR_START = 2006
telecom_ied = {}
for row in ws2.iter_rows(min_row=5, max_col=80, values_only=False):
    cell_a = str(row[0].value) if row[0].value else ''
    if cell_a.startswith('517 '):
        for yr in range(2013, 2025):
            if yr < 2024:
                idx = (yr - YEAR_START) * 4 + 3   # Q4 = anual
            else:
                idx = (yr - YEAR_START) * 4 + 1   # Q2 = enero-junio 2024
            v = row[idx + 1].value  # +1 porque row[0] es el label
            if v is not None and str(v) != 'C':
                telecom_ied[yr] = float(v)
            else:
                telecom_ied[yr] = 0.0
        break
wb2.close()

# ── 2. Preparar arrays ────────────────────────────────────────────────────────
years = list(range(2013, 2025))
ied_mexico = np.array([total_ied[y] for y in years])
ied_telecom = np.array([telecom_ied[y] for y in years])

# ── 3. Rango de eje X que usaría la gráfica (referencia) ─────────────────────
x_min = min(ied_telecom.min(), 0) - 4000
x_max = ied_mexico.max() + 5000

# ── 4. Imprimir resultados usados para la gráfica ─────────────────────────────
print(f"{'Año':<6} {'IED México':>16} {'IED Telecom':>16} {'% Telecom/México':>18}")
print("-" * 60)
for i, yr in enumerate(years):
    mx = ied_mexico[i]
    tc = ied_telecom[i]
    pct = (tc / mx * 100) if mx else float('nan')
    print(f"{yr:<6} {mx:>16,.2f} {tc:>16,.2f} {pct:>17.2f}%")

print("-" * 60)
print(f"{'Total':<6} {ied_mexico.sum():>16,.2f} {ied_telecom.sum():>16,.2f}")
print(f"{'Promedio':<6} {ied_mexico.mean():>16,.2f} {ied_telecom.mean():>16,.2f}")

print("\nEtiquetas de valor que se colocarían sobre cada barra:")
for i, yr in enumerate(years):
    mx_label = f"{ied_mexico[i]:,.0f}"
    tc_label = f"{ied_telecom[i]:,.2f}"
    print(f"  {yr}: México = {mx_label}   Telecom = {tc_label}")

print(f"\nRango del eje X (Millones de dólares): [{x_min:,.0f}, {x_max:,.0f}]")
