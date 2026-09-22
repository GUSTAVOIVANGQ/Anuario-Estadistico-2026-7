from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import anuario2026.pdf_export as pdf_export
from anuario2026.exports import ExportUnavailable, ensure_pdf
from anuario2026.pdf_export import ConverterBackend


def _transparent_png() -> BytesIO:
    stream = BytesIO()
    Image.new("RGBA", (1, 1), (0, 0, 0, 0)).save(stream, format="PNG")
    stream.seek(0)
    return stream


def _presentation(path: Path, slide_count: int = 3, *, with_a1: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width = Inches(10)
    deck.slide_height = Inches(5.625)
    for index in range(slide_count):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.75), Inches(0.7), Inches(8.5), Inches(1))
        box.text_frame.text = f"Diapositiva {index + 1}"
        if with_a1 and index == 0:
            helper = slide.shapes.add_textbox(Inches(2), Inches(1.5), Inches(6), Inches(3))
            helper.name = "ANUARIO_TEXT_LAYER_A_1"
            helper.text_frame.text = "Texto seleccionable de la figura A.1\nII\nPIB nacional"
            picture = slide.shapes.add_picture(
                _transparent_png(), Inches(2), Inches(1.5), Inches(6), Inches(3)
            )
            picture.name = "ANUARIO_IMAGE_A_1"
    deck.save(path)


def _simple_svg(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        """<svg xmlns="http://www.w3.org/2000/svg" width="600" height="300" viewBox="0 0 600 300">
<rect x="20" y="20" width="560" height="260" fill="none" stroke="#123456"/>
<text x="40" y="90" font-family="sans-serif" font-size="28">II</text>
<text x="40" y="150" font-family="sans-serif" font-size="28">PIB nacional</text>
<text x="40" y="210" font-family="sans-serif" font-size="20">Fuente: INEGI</text>
</svg>""",
        encoding="utf-8",
    )


def _converted_pdf(
    path: Path,
    page_count: int,
    *,
    text_by_page: dict[int, str] | None = None,
) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    for page_index in range(page_count):
        page = writer.add_blank_page(width=720, height=405)
        text = (text_by_page or {}).get(page_index)
        stream = DecodedStreamObject()
        commands = b"q 1 0 0 1 0 0 cm Q"
        if text:
            safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands += f" BT /F1 10 Tf 10 10 Td ({safe_text}) Tj ET".encode("ascii")
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): font_reference}
                    )
                }
            )
        stream.set_data(commands)
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_metadata({"/Producer": "Motor de prueba"})
    with path.open("wb") as output:
        writer.write(output)


def _project_with_pptx(
    tmp_path: Path,
    run_id: str = "corrida_pdf",
    *,
    with_a1: bool = False,
) -> tuple[Path, str]:
    (tmp_path / "reportes" / run_id).mkdir(parents=True)
    pptx = tmp_path / "entrega" / f"anuario_estadistico_2026_{run_id}.pptx"
    _presentation(pptx, with_a1=with_a1)
    if with_a1:
        svg = tmp_path / "build" / "figures" / "A" / "figura_a_1.svg"
        _simple_svg(svg)
        pptx.with_name(pptx.stem + "_ensamblaje.json").write_text(
            json.dumps(
                {
                    "inserted_count": 1,
                    "inserted": [
                        {
                            "figure_id": "A.1",
                            "slide_number": 1,
                            "svg_path": "build/figures/A/figura_a_1.svg",
                            "searchable_text_marker": "Texto seleccionable de la figura A.1",
                            "svg_texts": ["II", "PIB nacional", "Fuente: INEGI"],
                        }
                    ],
                    "embedded_svg_count": 1,
                    "transparent_svg_count": 1,
                    "fully_vector_svg_count": 1,
                    "embedded_raster_image_count": 0,
                    "raster_fallback_count": 0,
                    "searchable_text_layer_count": 1,
                    "svg_missing_count": 0,
                }
            ),
            encoding="utf-8",
        )
    return pptx, run_id


def test_pdf_uses_pptx_only_as_base_and_overlays_original_svg(tmp_path: Path, monkeypatch) -> None:
    pptx, run_id = _project_with_pptx(tmp_path, with_a1=True)
    calls: list[Path] = []

    monkeypatch.setattr(
        pdf_export,
        "_available_backends",
        lambda preferred=None: [ConverterBackend("Motor de prueba")],
    )

    def convert(backend, source, target, workspace):
        assert backend.name == "Motor de prueba"
        calls.append(source)
        base = Presentation(str(source))
        names = [shape.name for slide in base.slides for shape in slide.shapes]
        assert not any(name.startswith("ANUARIO_IMAGE_") for name in names)
        assert not any(name.startswith("ANUARIO_TEXT_LAYER_") for name in names)
        _converted_pdf(target, page_count=3)

    monkeypatch.setattr(pdf_export, "_run_converter", convert)

    pdf_path = ensure_pdf(tmp_path, run_id)
    reader = PdfReader(str(pdf_path))
    sidecar = json.loads(
        pdf_path.with_name(pdf_path.stem + "_pdf.json").read_text(encoding="utf-8")
    )
    page_text = reader.pages[0].extract_text() or ""

    assert len(calls) == 1
    assert calls[0] != pptx.resolve()
    assert calls[0].name == "presentacion_base_sin_figuras.pptx"
    assert len(reader.pages) == 3
    assert "II" in page_text
    assert "PIB nacional" in page_text
    assert "Fuente: INEGI" in page_text
    assert "Texto seleccionable de la figura A.1" not in page_text
    assert reader.metadata.title == "Anuario Estadístico 2026 - Presentación completa"
    assert reader.metadata.author == "Comisión Reguladora de Telecomunicaciones"
    assert sidecar["pdf_export_schema"] == 6
    assert sidecar["pdf_composition_mode"] == "pptx_base_plus_original_svg_overlay"
    assert sidecar["figures_converted_from_pptx"] is False
    assert sidecar["directly_composed_original_svg_count"] == 1
    assert sidecar["removed_figure_shapes_from_pdf_base"] == 1
    assert sidecar["removed_helper_text_layers_from_pdf_base"] == 1
    assert sidecar["validation"]["pdf_figures_are_direct_original_svg"] is True
    assert sidecar["validation"]["no_new_svg_generated_for_pdf"] is True
    assert sidecar["validation"]["all_svg_text_is_copyable"] is True


def test_pdf_cache_is_tied_to_pptx_and_svg_hash(tmp_path: Path, monkeypatch) -> None:
    _, run_id = _project_with_pptx(tmp_path, with_a1=True)
    conversion_count = 0
    monkeypatch.setattr(
        pdf_export,
        "_available_backends",
        lambda preferred=None: [ConverterBackend("Motor de prueba")],
    )

    def convert(backend, source, target, workspace):
        nonlocal conversion_count
        conversion_count += 1
        _converted_pdf(target, page_count=3)

    monkeypatch.setattr(pdf_export, "_run_converter", convert)

    first = ensure_pdf(tmp_path, run_id)
    second = ensure_pdf(tmp_path, run_id)
    assert first == second
    assert conversion_count == 1

    svg = tmp_path / "build" / "figures" / "A" / "figura_a_1.svg"
    svg.write_text(svg.read_text(encoding="utf-8").replace("Fuente: INEGI", "Fuente: CRT"), encoding="utf-8")
    # Al cambiar el original después del ensamblaje y no existir un SVG embebido
    # de control en este fixture, cambia la huella del bundle y se recompone.
    with pytest.raises(ExportUnavailable, match="perdió 1 textos SVG copiables"):
        ensure_pdf(tmp_path, run_id)


def test_pdf_rejects_a_base_conversion_with_missing_pages(tmp_path: Path, monkeypatch) -> None:
    _, run_id = _project_with_pptx(tmp_path)
    monkeypatch.setattr(
        pdf_export,
        "_available_backends",
        lambda preferred=None: [ConverterBackend("Motor de prueba")],
    )
    monkeypatch.setattr(
        pdf_export,
        "_run_converter",
        lambda backend, source, target, workspace: _converted_pdf(target, page_count=2),
    )

    with pytest.raises(ExportUnavailable, match="3 diapositivas.*2 páginas"):
        ensure_pdf(tmp_path, run_id)


def test_pdf_rejects_missing_original_svg(tmp_path: Path, monkeypatch) -> None:
    pptx, run_id = _project_with_pptx(tmp_path, with_a1=True)
    (tmp_path / "build" / "figures" / "A" / "figura_a_1.svg").unlink()
    monkeypatch.setattr(
        pdf_export,
        "_available_backends",
        lambda preferred=None: [ConverterBackend("Motor de prueba")],
    )

    with pytest.raises(ExportUnavailable, match="no existe el SVG original"):
        ensure_pdf(tmp_path, run_id)


def test_pdf_explains_when_no_office_renderer_is_installed(tmp_path: Path, monkeypatch) -> None:
    _, run_id = _project_with_pptx(tmp_path)
    monkeypatch.setattr(pdf_export, "_available_backends", lambda preferred=None: [])

    with pytest.raises(ExportUnavailable, match="PowerPoint.*LibreOffice"):
        ensure_pdf(tmp_path, run_id)
