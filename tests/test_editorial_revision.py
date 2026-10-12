from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from anuario2026.editorial_revision import (
    REGISTRY, apply_publication_revision, load_publication_revision, revision_fingerprint,
)
from anuario2026.exports import ensure_pptx
from anuario2026.pdf_export import _add_delivery_outline
from anuario2026.presentation_headers import (
    HEADER_DESIGN_VERSION, apply_uniform_headers, validate_uniform_headers,
)


ROOT = Path(__file__).resolve().parents[1]


def test_reviewed_copy_removal_and_page_mapping(tmp_path: Path):
    deck = Presentation(ROOT / "assets/presentation/anuario_estadistico_2026_automatizable.pptx")
    figure = deck.slides[10]
    covers = [deck.slides[i]._element.xml for i in (0, 130)]
    dimensions = deck.slide_width, deck.slide_height
    apply_uniform_headers(deck)
    revision = load_publication_revision(ROOT)
    report = apply_publication_revision(ROOT, deck)
    assert len(deck.slides) == 130
    assert deck.slides[9] is figure
    assert report["template_to_publication_pages"]["11"] == 10
    assert report["template_to_publication_pages"]["131"] == 130
    assert [deck.slides[i]._element.xml for i in (0, 129)] == covers
    assert (deck.slide_width, deck.slide_height) == dimensions
    legal = "\n".join(s.text for s in deck.slides[2].shapes if s.has_text_frame)
    for section in revision["legal_sections"].values():
        assert " ".join(section["body"].replace("*", "").split()) in " ".join(legal.split())
    glossary = "\n".join(s.text for s in deck.slides[3].shapes if s.has_text_frame)
    assert revision["gemini_definition"] in glossary
    assert "DGP: Dirección General de Planeación" in glossary
    for slide in deck.slides:
        text = " ".join(" ".join(s.text for s in slide.shapes if s.has_text_frame).split())
        assert "CGPE" not in text
        assert "Coordinación General de Planeación Estratégica" not in text
        assert "BARD DE GOOGLE" not in text
    assert "ANUARIO_GLOSARIO_DESIGN_2026" in [s.name for s in deck.slides[3].shapes]
    path = tmp_path / "revised.pptx"
    deck.save(path)
    reopened = Presentation(path)
    validate_uniform_headers(reopened)
    assert len(reopened.slides) == 130
    assert reopened.slides[9].part.partname == "/ppt/slides/slide10.xml"
    for number, slide in enumerate(reopened.slides, 1):
        for shape in slide.shapes:
            if shape.name == "ANUARIO_PAGE_NUMBER":
                assert shape.text == str(number)
    assert revision["gemini_source"] in reopened.slides[3].notes_slide.notes_text_frame.text


def test_export_cache_detects_editorial_changes(tmp_path: Path, monkeypatch):
    run_id = "revision"
    (tmp_path / "reportes" / run_id).mkdir(parents=True)
    (tmp_path / "config").mkdir()
    (tmp_path / "config/proyecto.json").write_text("{}")
    registry = tmp_path / REGISTRY
    registry.parent.mkdir(parents=True)
    registry.write_text(json.dumps({"publication_revision": {"version": 1}}))
    output = tmp_path / "entrega" / f"anuario_estadistico_2026_{run_id}.pptx"
    output.parent.mkdir()
    output.write_bytes(b"old")
    sidecar = output.with_name(output.stem + "_ensamblaje.json")
    sidecar.write_text(json.dumps({"header_design": {"version": HEADER_DESIGN_VERSION}}))
    calls = []

    def assemble(root, run_dir, *, output):
        calls.append(run_dir.name)
        output.write_bytes(b"revised")
        sidecar.write_text(json.dumps({
            "header_design": {"version": HEADER_DESIGN_VERSION},
            "publication_revision": {"fingerprint": revision_fingerprint(root)},
        }))

    monkeypatch.setattr("anuario2026.exports.assemble_pptx", assemble)
    ensure_pptx(tmp_path, run_id)
    ensure_pptx(tmp_path, run_id)
    assert calls == [run_id]
    registry.write_text(json.dumps({"publication_revision": {"version": 2}}))
    ensure_pptx(tmp_path, run_id)
    assert calls == [run_id, run_id]


def test_pdf_outline_uses_publication_pages(tmp_path: Path, monkeypatch):
    monkeypatch.setattr("anuario2026.pdf_export._reference_index_entries",
                        lambda root: [("GLOSARIO", 5, 0), ("A. INDICADORES ECONÓMICOS", 11, 0)])
    writer = PdfWriter()
    for _ in range(12):
        writer.add_blank_page(width=720, height=405)
    _add_delivery_outline(writer, tmp_path, 12, {"5": 4, "11": 10})
    path = tmp_path / "outline.pdf"
    writer.write(path)
    reader = PdfReader(path)
    assert reader.get_destination_page_number(reader.outline[2]) == 3
    assert reader.get_destination_page_number(reader.outline[3]) == 9
