from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from anuario2026.exports import ensure_pptx
from anuario2026.presentation_headers import (
    BIT_URL,
    HEADER_DESIGN_VERSION,
    HEADER_NAMES,
    _is_header,
    apply_uniform_headers,
    validate_uniform_headers,
)


def test_uniform_headers_preserve_body_covers_and_reference(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    deck = Presentation(root / "assets/presentation/anuario_estadistico_2026_automatizable.pptx")
    unchanged = {n: deck.slides[n - 1]._element.xml for n in (1, 7, 131)}
    body = {
        n: [s._element.xml for s in deck.slides[n - 1].shapes if not _is_header(s)]
        for n in range(2, 131)
    }
    reference_images = [s.image.blob for s in deck.slides[6].shapes if _is_header(s) and s.shape_type == 13]
    apply_uniform_headers(deck)
    for n, xml in unchanged.items():
        assert deck.slides[n - 1]._element.xml == xml
    for n in range(2, 131):
        slide = deck.slides[n - 1]
        assert [s._element.xml for s in slide.shapes if not _is_header(s)] == body[n]
        images = [s.image.blob for s in slide.shapes if _is_header(s) and s.shape_type == 13]
        assert images == reference_images
        if n == 7:
            continue
        top = [s for s in slide.shapes if _is_header(s)]
        assert tuple(s.name for s in top) == HEADER_NAMES
        url = top[-1]
        run = url._element.find(".//" + qn("a:rPr"))
        assert run.find(qn("a:hlinkClick")) is None
        assert run.get("u") == "sng"
        assert run.find(qn("a:solidFill") + "/" + qn("a:srgbClr")).get("val") == "FFFFFF"
        link = url._element.find(".//" + qn("p:cNvPr") + "/" + qn("a:hlinkClick"))
        assert slide.part.rels[link.get(qn("r:id"))].target_ref == BIT_URL
    # Guardar y volver a abrir verifica relaciones de imágenes y enlaces.
    path = tmp_path / "headers.pptx"
    deck.save(path)
    reopened = Presentation(path)
    validate_uniform_headers(reopened)
    apply_uniform_headers(reopened)
    validate_uniform_headers(reopened)
    assert len(reopened.slides[10].shapes) == len(deck.slides[10].shapes)


def test_existing_run_rebuilds_stale_pptx_only_once(tmp_path: Path, monkeypatch) -> None:
    run_id = "corrida_anterior"
    (tmp_path / "reportes" / run_id).mkdir(parents=True)
    (tmp_path / "config").mkdir()
    (tmp_path / "config/proyecto.json").write_text("{}", encoding="utf-8")
    output = tmp_path / "entrega" / f"anuario_estadistico_2026_{run_id}.pptx"
    output.parent.mkdir()
    output.write_bytes(b"pptx anterior sin barra")
    report = output.with_name(output.stem + "_ensamblaje.json")
    report.write_text("{}", encoding="utf-8")
    calls = []

    def assemble(root, run_dir, *, output):
        calls.append(run_dir.name)
        output.write_bytes(b"pptx con barra aprobada")
        report.write_text(json.dumps({
            "header_design": {"version": HEADER_DESIGN_VERSION},
        }), encoding="utf-8")

    monkeypatch.setattr("anuario2026.exports.assemble_pptx", assemble)
    assert ensure_pptx(tmp_path, run_id) == output
    assert output.read_bytes() == b"pptx con barra aprobada"
    assert ensure_pptx(tmp_path, run_id) == output
    assert calls == [run_id]
