from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_PATH = PROJECT_ROOT / "scripts" / "figures" / "figura_a_5.py"
SPEC = importlib.util.spec_from_file_location("figura_a_5_test", SCRIPT_PATH)
assert SPEC and SPEC.loader
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


def _total_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Originales y actualización"
    worksheet.cell(2, 1, "Año")
    worksheet.cell(2, 2, "Periodo")
    worksheet.cell(2, 4, "Datos actualizados al segundo trimestre 2026")
    row = 3
    for year in range(2013, 2027):
        worksheet.cell(row, 1, year)
        worksheet.cell(row, 2, "Enero - marzo" if year == 2026 else "Enero - diciembre")
        worksheet.cell(row, 4, year * 10.0)
        row += 1
    workbook.save(path)


def _activity_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Por Actividad Económica"
    worksheet.cell(5, 1, "517 Telecomunicaciones")
    column = 2
    for year in range(2013, 2027):
        for quarter in range(1, 5):
            if quarter == 1:
                worksheet.cell(3, column, year)
            worksheet.cell(4, column, quarter)
            worksheet.cell(5, column, year + quarter / 10)
            column += 1
    workbook.save(path)


def test_a5_reads_comparable_periods_from_manual_workbooks(tmp_path: Path):
    total_path = tmp_path / MODULE.TOTAL_FILENAME
    activity_path = tmp_path / MODULE.ACTIVITY_FILENAME
    _total_workbook(total_path)
    _activity_workbook(activity_path)

    total = MODULE.read_total_ied(total_path)
    telecom = MODULE.read_telecom_ied(activity_path)
    result = MODULE.calculate_series(total, telecom)

    assert result["anio"].tolist() == list(range(2013, 2027))
    assert total[(2025, 4)] == 20250.0
    assert telecom[(2025, 4)] == 2025.4
    assert telecom[(2026, 1)] == 2026.1
    assert result.iloc[-1]["periodo"] == "enero-marzo"
    assert result.iloc[-1]["participacion_telecom_pct"] == pytest.approx(
        telecom[(2026, 1)] / total[(2026, 1)] * 100
    )


def test_a5_reports_exact_manual_input_location_when_a_file_is_missing(tmp_path: Path):
    total_path = tmp_path / MODULE.TOTAL_FILENAME
    total_path.touch()

    with pytest.raises(FileNotFoundError, match="El programa no descarga estos insumos"):
        MODULE._input_paths((total_path,), tmp_path)
