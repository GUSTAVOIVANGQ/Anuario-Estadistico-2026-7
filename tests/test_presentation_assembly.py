from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from pptx import Presentation

from anuario2026.presentation import TRANSPARENT_BOOTSTRAP_PNG, assemble_from_template


def _project_root() -> Path:
    return Path(__file__).resolve().parents[1]


def test_template_manifest_matches_placeholders() -> None:
    root = _project_root()
    config = json.loads((root / "config" / "proyecto.json").read_text(encoding="utf-8"))
    template = root / config["presentation"]["template"]
    manifest = json.loads(
        (root / config["presentation"]["manifest"]).read_text(encoding="utf-8")
    )
    presentation = Presentation(str(template))

    assert len(presentation.slides) == manifest["slide_count"] == 119
    assert len(manifest["entries"]) == manifest["figure_count"] == 105

    for entry in manifest["entries"]:
        slide = presentation.slides[entry["slide_number"] - 1]
        names = {shape.name for shape in slide.shapes}
        assert entry["figure_shape_name"] in names
        assert entry["narrative_shape_name"] in names


def test_assemble_inserts_image_and_clears_token(tmp_path: Path) -> None:
    source_root = _project_root()
    project_root = tmp_path / "project"
    (project_root / "config").mkdir(parents=True)
    (project_root / "assets" / "presentation").mkdir(parents=True)
    (project_root / "build" / "figures" / "A").mkdir(parents=True)

    source_config = json.loads(
        (source_root / "config" / "proyecto.json").read_text(encoding="utf-8")
    )
    (project_root / "config" / "proyecto.json").write_text(
        json.dumps(source_config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    template_rel = Path(source_config["presentation"]["template"])
    manifest_rel = Path(source_config["presentation"]["manifest"])
    shutil.copy2(source_root / template_rel, project_root / template_rel)
    shutil.copy2(
        source_root / "assets" / "presentation" / "narrativas_figuras_2026.json",
        project_root / "assets" / "presentation" / "narrativas_figuras_2026.json",
    )

    manifest = json.loads((source_root / manifest_rel).read_text(encoding="utf-8"))
    manifest["entries"] = [entry for entry in manifest["entries"] if entry["figure_id"] == "A.1"]
    manifest["figure_count"] = 1
    (project_root / manifest_rel).write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    source_svg = source_root / "build" / "figures" / "A" / "figura_a_1.svg"
    target_svg = project_root / "build" / "figures" / "A" / "figura_a_1.svg"
    svg = source_svg.read_text(encoding="utf-8").replace(
        'style="fill: #ffffff"',
        'style="fill: none"',
        1,
    )
    target_svg.write_text(svg, encoding="utf-8")

    output = project_root / "entrega" / "test.pptx"
    result = assemble_from_template(project_root, output, strict=True)
    assert result.inserted == 1
    assert result.missing == ()
    assert result.errors == ()

    presentation = Presentation(str(output))
    slide = presentation.slides[5]
    by_name = {shape.name: shape for shape in slide.shapes}
    assert by_name["ANUARIO_FIGURE_A_1"].text == ""
    assert "ANUARIO_IMAGE_A_1" in by_name
    assert "ANUARIO_TEXT_LAYER_A_1" in by_name
    assert "Texto seleccionable de la figura A.1" in by_name["ANUARIO_TEXT_LAYER_A_1"].text
    assert "PIB nacional" in by_name["ANUARIO_TEXT_LAYER_A_1"].text
    assert "En el segundo trimestre de 2026" in by_name["ANUARIO_TEXT_A_1"].text
    assert "$ 26,144.13" in by_name["ANUARIO_TEXT_A_1"].text

    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert report["embedded_svg_count"] == 1
    assert report["searchable_text_layer_count"] == 1
    assert report["svg_missing_count"] == 0
    assert report["transparent_svg_count"] == 1
    assert report["raster_fallback_count"] == 0
    assert report["narrative_inserted_count"] == 1
    assert report["narrative_missing_count"] == 0
    assert report["inserted"][0]["svg_text_elements"] > 0
    assert "PIB nacional" in " ".join(report["inserted"][0]["svg_texts"])

    with zipfile.ZipFile(output) as archive:
        content_types = archive.read("[Content_Types].xml")
        assert (
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            in content_types
        )
        assert b"ns0:Types" not in content_types
        assert "ppt/media/anuario_figure_a_1.svg" in archive.namelist()
        assert archive.read("ppt/media/anuario_figure_a_1.svg") == (
            project_root / "build" / "figures" / "A" / "figura_a_1.svg"
        ).read_bytes()
        slide_xml = archive.read("ppt/slides/slide6.xml")
        assert b"svgBlip" not in slide_xml
        relationships = archive.read("ppt/slides/_rels/slide6.xml.rels")
        assert b"../media/anuario_figure_a_1.svg" in relationships
        assert all(
            archive.read(name) != TRANSPARENT_BOOTSTRAP_PNG
            for name in archive.namelist()
            if name.startswith("ppt/media/")
        )
