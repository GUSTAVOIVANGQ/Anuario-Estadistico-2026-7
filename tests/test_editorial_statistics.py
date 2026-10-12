from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

from anuario2026.editorial_revision import apply_publication_revision, load_publication_revision
from anuario2026.editorial_statistics import annual_income, calculate_highlights, quarterly_gdp


ROOT = Path(__file__).resolve().parents[1]


def income_frame():
    return pd.DataFrame({"anio": [2024] * 4 + [2025] * 4,
                         "trimestre": [1, 2, 3, 4] * 2,
                         "ingresos_miles_millones_pesos": [10, 20, 30, 40, 20, 40, 60, 80]})


def test_income_uses_complete_annual_sum_and_rejects_missing_or_duplicate_quarters():
    frame = income_frame()
    result = annual_income(frame, 2025)
    assert result["value_mmdp"] == 200
    assert result["previous_mmdp"] == 100
    assert result["annual_growth_pct"] == 100
    with pytest.raises(ValueError, match="Faltan trimestres"):
        annual_income(frame.iloc[:-1], 2025)
    with pytest.raises(ValueError, match="duplicados"):
        annual_income(pd.concat([frame, frame.iloc[:1]]), 2025)


def test_gdp_compares_fourth_quarters_in_same_units():
    frame = pd.DataFrame({"anio": [2024, 2025, 2025], "trimestre": [4, 2, 4],
                          "tyr_miles_millones_pesos": [400, 700, 440]})
    result = quarterly_gdp(frame, 2025)
    assert result["value_mmdp"] == 440
    assert result["annual_growth_pct"] == pytest.approx(10)


def test_complete_run_data_supersedes_verified_editorial_snapshot(tmp_path):
    (tmp_path / "datos_usados").mkdir()
    path = tmp_path / "datos_usados/a_6_datos_usados.csv"
    income_frame().to_csv(path, index=False)
    result = calculate_highlights(ROOT, load_publication_revision(ROOT), tmp_path)
    assert result["income"]["value_mmdp"] == 200
    assert result["income"]["data_file"] == str(path)
    income_frame().iloc[:-1].to_csv(path, index=False)
    fallback = calculate_highlights(ROOT, load_publication_revision(ROOT), tmp_path)
    assert fallback["income"]["value_mmdp"] == pytest.approx(692.20790265444)
    assert fallback["income"]["data_file"] != str(path)


def test_revision_two_texts_and_figure_design_survive_assembly():
    deck = Presentation(ROOT / "assets/presentation/anuario_estadistico_2026_automatizable.pptx")
    unchanged = _xml_body(deck.slides[8], {"TextBox 14", "TextBox 17", "TextBox 8"})
    settings = load_publication_revision(ROOT)
    report = apply_publication_revision(ROOT, deck)
    normalize = lambda value: " ".join(value.replace("**", "").split())
    intro = normalize(" ".join(s.text for s in deck.slides[5].shapes if s.has_text_frame))
    assert normalize(settings["introduction"]["body"]) in intro
    notes = normalize(" ".join(s.text for s in deck.slides[6].shapes if s.has_text_frame))
    for item in settings["introduction_notes"]:
        assert normalize(item["body"]) in notes
    assert "IFT" not in intro
    assert "IFT" not in notes
    link = next(s for s in deck.slides[6].shapes if s.name == "INTRO8_NOTA_5_1")
    assert link.click_action.hyperlink.address == "https://bit.crt.gob.mx/BitWebApp/PortalBIT.xhtml?view="
    url_run = next(r for r in link.text_frame.paragraphs[0].runs if r.text.startswith("https://"))
    assert str(url_run.font.color.rgb) == "296B8C"
    assert url_run.hyperlink.address is None
    points = normalize(" ".join(s.text for s in deck.slides[7].shapes if s.has_text_frame))
    for value in report["economic_highlights"].values():
        assert normalize(value["text"]) in points
    assert _xml_body(deck.slides[7], {"TextBox 14", "TextBox 17", "TextBox 8"}) == unchanged
    assert next(s for s in deck.slides[7].shapes if s.name == "TextBox 8").text == "8"


def _xml_body(slide, excluded):
    return [s._element.xml for s in slide.shapes if s.name not in excluded]
