"""Adopta un PPTX de 131 páginas y sincroniza sus marcadores con el generador.

Uso: python scripts/preparar_plantilla_estructura_2024.py RUTA_DEL_BORRADOR.pptx
"""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets" / "presentation"
TEMPLATE = ASSETS / "anuario_estadistico_2026_automatizable.pptx"
MANIFEST = ASSETS / "anuario_estadistico_2026_manifest.json"
CSV_MANIFEST = ASSETS / "anuario_estadistico_2026_manifest.csv"
EDITORIAL = ASSETS / "textos_editoriales_2026.json"
TOKEN = re.compile(r"\{\{(TEXTO|TABLA|CONTRAPORTADA):([^}]+)\}\}")


def prepare(source: Path) -> None:
    deck = Presentation(str(source))
    if len(deck.slides) != 131:
        raise ValueError(f"Se esperaban 131 páginas y hay {len(deck.slides)}")

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    locations: dict[str, list[tuple[int, object]]] = {}
    editorial_slots = []
    for page, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            locations.setdefault(shape.name, []).append((page, shape))
            if not shape.has_text_frame:
                continue
            if shape.name.startswith("ANUARIO_TEXT_"):
                # La narrativa de las figuras se gestiona con su propio registro.
                continue
            for match in TOKEN.finditer(shape.text):
                editorial_slots.append(
                    {
                        "key": match.group(2),
                        "kind": match.group(1).lower(),
                        "slide_number": page,
                        "shape_name": shape.name,
                        "token": match.group(0),
                    }
                )

    for entry in manifest["entries"]:
        figure = locations.get(entry["figure_shape_name"], [])
        narrative = locations.get(entry["narrative_shape_name"], [])
        if len(figure) != 1 or len(narrative) != 1 or figure[0][0] != narrative[0][0]:
            raise ValueError(f"Marcadores incompletos o duplicados: {entry['figure_id']}")
        page, shape = figure[0]
        if entry["token"] not in shape.text:
            raise ValueError(f"Token incorrecto en {entry['figure_id']}")
        entry["slide_number"] = page
        entry["figure_bounds_inches"] = {
            key: round(getattr(shape, key) / 914400, 4)
            for key in ("left", "top", "width", "height")
        }

    # El índice conserva también las erratas de la edición 2024: la réplica
    # solicitada es literal, mientras que el cuerpo editorial sí se revisa.

    ids = [item["key"] for item in editorial_slots]
    if len(ids) != len(set(ids)):
        raise ValueError("Hay claves editoriales duplicadas")
    manifest["slide_count"] = len(deck.slides)
    manifest["editorial_slots"] = editorial_slots
    deck.save(str(TEMPLATE))
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    fields = [
        "slide_number", "section", "figure_id", "title", "token",
        "figure_shape_name", "narrative_shape_name", "expected_image",
        "left", "top", "width", "height",
    ]
    with CSV_MANIFEST.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for entry in manifest["entries"]:
            bounds = entry["figure_bounds_inches"]
            writer.writerow({
                **{field: entry.get(field, "") for field in fields},
                **bounds,
            })

    # Conservar la redacción ya revisada al regenerar; nunca inventar cifras.
    previous = json.loads(EDITORIAL.read_text(encoding="utf-8")).get("slots", {}) if EDITORIAL.exists() else {}
    EDITORIAL.write_text(
        json.dumps({"slots": {key: previous.get(key) for key in ids}}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Plantilla: {TEMPLATE} ({len(deck.slides)} diapositivas)")
    print(f"Figuras: {len(manifest['entries'])}; campos editoriales: {len(editorial_slots)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Uso: python scripts/preparar_plantilla_estructura_2024.py BORRADOR.pptx")
    prepare(Path(sys.argv[1]))
