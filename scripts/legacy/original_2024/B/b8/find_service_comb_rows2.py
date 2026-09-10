import glob
from openpyxl import load_workbook

SERVICE_TOKENS = ["telefon", "telefono", "telefonia", "telefonía", "tv", "television", "televisión", "internet", "restringida", "fija"]

files = sorted(glob.glob(r"..\\datos\\B.1\\*.xlsx", recursive=False))

for f in files:
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
    except Exception as e:
        print(f"ERROR opening {f}: {e}")
        continue
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            try:
                line = " | ".join([str(c).strip().lower() for c in row if c is not None])
            except Exception:
                continue
            tokens_found = sum(1 for t in SERVICE_TOKENS if t in line)
            if tokens_found >= 2:
                print(f"File: {f} | Sheet: {sheet} | Row: {i} | tokens={tokens_found}")
                print(line)
                for j in range(i+1, i+7):
                    try:
                        r2 = ws[j]
                    except Exception:
                        break
                    values = [str(c.value).strip() for c in r2 if c.value is not None]
                    if values:
                        print("-> "+" | ".join(values))
                print("="*80)
    wb.close()

print('Done')
