import glob
from openpyxl import load_workbook
import re

KEYWORDS = [
    "tres servicios",
    "dos servicios",
    "un servicio",
    "ninguno",
    "combinacion",
    "combinación",
    "combinaciones",
    "telefonía fija",
    "telefonia fija",
    "televisión restringida",
    "television restringida",
    "tv restringida",
    "internet",
    "combinación de servicios",
]

pattern = re.compile("|".join([re.escape(k) for k in KEYWORDS]), re.IGNORECASE)

files = sorted(glob.glob(r"..\\datos\\B.1\\*.xlsx", recursive=False))

for f in files:
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
    except Exception as e:
        print(f"ERROR opening {f}: {e}")
        continue
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            try:
                line = " | ".join([str(c).strip() for c in row if c is not None])
            except Exception:
                continue
            if pattern.search(line):
                print(f"File: {f} | Sheet: {sheet} | Row: {i}")
                print(line)
                # print next 6 rows for context
                ctx = []
                for j in range(i+1, i+7):
                    try:
                        r2 = ws[j]
                    except Exception:
                        break
                    values = [str(c.value).strip() for c in r2 if c.value is not None]
                    if values:
                        print("-> "+" | ".join(values))
                print("="*80)
    wb.close()

print('Done')
